from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles; styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10.5)
for s,size in [('Title',24),('Heading 1',16),('Heading 2',13)]: styles[s].font.name='Microsoft YaHei'; styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles[s].font.size=Pt(size); styles[s].font.bold=True
doc.add_heading('知源 CloudBase 演示部署方案',0); p=doc.add_paragraph('适用于黑客松评委体验的轻量 Demo'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('一 结论',1)
doc.add_paragraph('本项目没有 MySQL、MongoDB、Redis 等外部数据库。当前代码使用 Python FastAPI、静态前端和 Dockerfile。演示模式下，笔记、卡片和历史只用于当前浏览器会话，不要求配置数据库。后端仍可能产生少量容器本地文件（匿名身份、复盘或上传临时文件），容器重启后不保证保留，因此无需购买或绑定持久化数据库。')
doc.add_paragraph('推荐使用 CloudBase 云托管，从 GitHub main 分支按 Dockerfile 构建。知乎 OAuth 暂不作为首场演示必需项；知乎搜索和研究模式可单独演示。')
doc.add_heading('二 项目是否有数据库写入',1)
for t in ['没有外部数据库连接：requirements.txt 和代码中没有 MySQL、PostgreSQL、MongoDB、Redis 客户端。','存在文件级写入：data/、uploads/ 可能保存运行时文件；这不是数据库，且 Demo 模式不应依赖其长期存在。','演示模式会清理浏览器本地的笔记、卡片和历史，刷新页面后重新开始，符合即插即用定位。','如果以后要做正式产品，再接入 CloudBase 数据库或对象存储，并关闭 DEMO_MODE。']: doc.add_paragraph(t, style='List Bullet')
doc.add_heading('三 CloudBase 部署步骤',1)
steps=['将仓库推送到 GitHub：https://github.com/Scarlett-yzy/isoLogic，确认分支为 main，且推的是最新提交。','进入 CloudBase 控制台 → 云托管 → 新建服务。代码来源选择 GitHub 构建，仓库选择上述地址，分支选择 main，构建方式选择 Dockerfile。','容器端口填写 8000；请求超时填写 300 秒；最小实例 1，最大实例 1；建议 1 核 2 GB。启动命令留空，由 Dockerfile 自动执行。','在环境变量中填写模型 API 和知乎搜索配置。不要把真实密钥写进 GitHub。','点击部署，等待镜像构建完成。首次构建可能需要数分钟。','打开部署后的域名首页和 /api/health 验收。']
for i,t in enumerate(steps,1): doc.add_paragraph(f'{i}. {t}')
doc.add_heading('四 必填环境变量',1)
table=doc.add_table(rows=1, cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
for c,v in zip(table.rows[0].cells,['变量','示例','用途']): c.text=v
rows=[('LC_API_KEY','交付方提供的模型密钥','AI 解释、研究整合、复盘'),('LC_BRIDGE','https://api.openai-next.com','模型网关'),('LC_MODEL','gpt-4o-mini','模型名称'),('LC_TIMEOUT','180','单次模型超时'),('LC_DISCOVER_TIMEOUT','240','发现同源超时'),('LC_THINKING','none','关闭额外思考参数'),('LC_SECRET_KEY','随机长字符串','匿名身份签名'),('ZHIHU_ACCESS_SECRET','知乎开放平台密钥','知乎搜索')]
for row in rows: cells=table.add_row().cells; [setattr(c,'text',v) for c,v in zip(cells,row)]
doc.add_heading('五 OAuth 说明',1)
doc.add_paragraph('OAuth 不是部署必需项。启用后，站点管理员配置一次 Client ID、Client Secret 和公网回调地址；每位用户仍需在知乎页面登录并授权自己的账号。令牌必须按用户隔离保存。若知乎授权页没有收藏或点赞读取权限，就不能实现对应同步。')
doc.add_paragraph('公网部署后回调地址示例： https://你的CloudBase域名/api/zhihu/oauth/callback。该地址必须与知乎后台登记值完全一致。')
doc.add_heading('六 验收清单',1)
for t in ['访问首页，页面能正常打开。','访问 /api/health，返回 code: 0 且 api: ok。','测试“知乎搜索”，确认能返回文章。','测试“研究模式”，确认能生成整合笔记和参考来源。','测试“发现同源”和“笔记复盘”。','确认刷新页面后演示数据清空，不把 Demo 当成永久资料库。']: doc.add_paragraph(t, style='List Bullet')
doc.add_heading('七 常见问题',1)
doc.add_paragraph('如果页面仍显示旧功能，执行 Ctrl + F5，并确认 CloudBase 构建的确实是 main 分支最新提交。如果 AI 返回 503，检查 LC_API_KEY、LC_BRIDGE 和模型网关连通性。如果知乎 OAuth 返回 503，说明尚未配置 ZHIHU_OAUTH_CLIENT_ID；这不影响知乎搜索和研究模式。')
doc.save('CloudBase演示部署方案.docx')
